from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from job_applier.models.job_application import ScanRun, UnknownQuestion
from job_applier.storage.excel_report_writer import ExcelReportWriter, SpreadsheetLockedError
from tests.storage.conftest import MOMENT, application

QUESTION = UnknownQuestion(
    question="Você tem CNH?",
    occurrences=2,
    last_company="Empresa X",
    last_title="Dev Júnior",
    first_seen_at=MOMENT,
    last_seen_at=datetime(2026, 10, 30, 9, 0),
)


@pytest.fixture
def writer(tmp_path) -> ExcelReportWriter:
    return ExcelReportWriter(tmp_path / "data" / "candidaturas.xlsx", project_root=tmp_path)


class TestExcelReportWriter:
    def test_writes_applications_sheet_with_headers_and_values(self, writer, tmp_path):
        writer.write([application()], [])

        sheet = load_workbook(tmp_path / "data" / "candidaturas.xlsx")["Vagas"]
        headers = [cell.value for cell in sheet[1]]
        row = {header: cell for header, cell in zip(headers, sheet[2], strict=True)}
        assert headers[:6] == ["Data", "Empresa", "Vaga", "Local", "Status", "Aderência"]
        assert row["Data"].value == MOMENT
        assert row["Status"].value == "CV gerado"
        assert row["Aderência"].value == 0.75 and row["Aderência"].number_format == "0%"
        assert row["Tem"].value == "Python, Docker"
        assert row["Vaga de suporte"].value == "Não"

    def test_links_resume_relative_to_the_spreadsheet_and_job_url(self, writer, tmp_path):
        writer.write([application()], [])

        sheet = load_workbook(tmp_path / "data" / "candidaturas.xlsx")["Vagas"]
        headers = [cell.value for cell in sheet[1]]
        row = {header: cell for header, cell in zip(headers, sheet[2], strict=True)}
        assert row["CV"].value == "2026-10-29_Empresa-X_Dev-Junior.pdf"
        assert row["CV"].hyperlink.target == "cvs/2026-10-29_Empresa-X_Dev-Junior.pdf"
        assert row["Link da vaga"].hyperlink.target == "https://www.linkedin.com/jobs/view/1/"

    def test_writes_unknown_questions_sheet(self, writer, tmp_path):
        writer.write([], [QUESTION])

        sheet = load_workbook(tmp_path / "data" / "candidaturas.xlsx")["Perguntas"]
        assert [cell.value for cell in sheet[2]][:3] == ["Você tem CNH?", 2, "Empresa X - Dev Júnior"]

    def test_freezes_header_and_enables_filters(self, writer, tmp_path):
        writer.write([application()], [])

        sheet = load_workbook(tmp_path / "data" / "candidaturas.xlsx")["Vagas"]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:R2"

    def test_shows_linkedin_details_and_skipped_jobs_without_resume(self, writer, tmp_path):
        skipped = application(
            resume_path=None,
            status="ignorada",
            notes="Vaga de suporte",
            apply_type="candidatura_simplificada",
            applicants=120,
            posted_on=date(2026, 10, 27),
        )

        writer.write([skipped], [])

        sheet = load_workbook(tmp_path / "data" / "candidaturas.xlsx")["Vagas"]
        headers = [cell.value for cell in sheet[1]]
        row = {header: cell for header, cell in zip(headers, sheet[2], strict=True)}
        assert row["Status"].value == "Ignorada"
        assert row["Observação"].value == "Vaga de suporte"
        assert row["Candidatura"].value == "Candidatura simplificada"
        assert row["Candidatos"].value == 120
        assert row["Publicada em"].value == datetime(2026, 10, 27)
        assert row["CV"].value is None and row["CV"].hyperlink is None

    def test_writes_scan_runs_sheet(self, writer, tmp_path):
        run = ScanRun(
            started_at=MOMENT,
            finished_at=datetime(2026, 10, 29, 11, 0),
            status="interrompida",
            reason="Captcha",
            jobs_seen=4,
            jobs_recorded=3,
        )

        writer.write([], [], [run])

        sheet = load_workbook(tmp_path / "data" / "candidaturas.xlsx")["Execuções"]
        assert [cell.value for cell in sheet[2]][2:] == ["Interrompida", "Captcha", 4, 3]

    def test_reports_locked_spreadsheet_and_leaves_no_temporary_file(self, writer, tmp_path):
        (tmp_path / "data" / "candidaturas.xlsx").mkdir(parents=True)

        with pytest.raises(SpreadsheetLockedError, match="aberta no Excel"):
            writer.write([application()], [])

        assert list((tmp_path / "data").glob("~*")) == []
