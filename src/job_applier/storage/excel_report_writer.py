import os
from collections.abc import Mapping, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from job_applier.models.job_application import (
    ApplicationStatus,
    JobApplication,
    ScanRun,
    ScanStatus,
    UnknownQuestion,
)
from job_applier.models.job_posting import ApplyType


class SpreadsheetLockedError(Exception):
    """A planilha não pôde ser gravada (normalmente porque está aberta no Excel)."""


class ExcelReportWriter:
    """Gera a planilha data/candidaturas.xlsx a partir do banco, com as abas Vagas, Perguntas e Execuções.

    A planilha é só um relatório: é recriada inteira a cada execução, então se uma gravação
    falhar os dados continuam no banco e a próxima execução a atualiza.
    """

    _APPLICATION_COLUMNS = {
        "Data": 17, "Empresa": 24, "Vaga": 42, "Local": 24, "Status": 20, "Aderência": 11, "Candidatura": 24,
        "Candidatos": 12, "Publicada em": 13, "Foco": 14, "Idioma": 8, "Vaga de suporte": 15, "Tem": 50,
        "Falta": 40, "Observação": 40, "CV": 50, "Link da vaga": 14, "Atualizado em": 17,
    }
    _QUESTION_COLUMNS = {"Pergunta": 70, "Ocorrências": 12, "Última vaga": 50, "Primeira vez": 17, "Última vez": 17}
    _RUN_COLUMNS = {"Início": 17, "Fim": 17, "Situação": 14, "Motivo": 70, "Vagas vistas": 13, "Vagas registradas": 17}

    _STATUS_LABELS = {
        ApplicationStatus.CV_GENERATED: "CV gerado",
        ApplicationStatus.APPLIED: "Candidatura enviada",
        ApplicationStatus.SKIPPED: "Ignorada",
        ApplicationStatus.FAILED: "Erro",
    }
    _APPLY_TYPE_LABELS = {
        ApplyType.EASY_APPLY: "Candidatura simplificada",
        ApplyType.EXTERNAL: "Site da empresa",
        ApplyType.CLOSED: "Encerrada",
        ApplyType.UNKNOWN: "Desconhecida",
    }
    _RUN_STATUS_LABELS = {
        ScanStatus.RUNNING: "Em andamento",
        ScanStatus.COMPLETED: "Concluída",
        ScanStatus.STOPPED: "Interrompida",
    }
    _DATE_TIME_FORMAT = "dd/mm/yyyy hh:mm"
    _DATE_FORMAT = "dd/mm/yyyy"
    _HEADER_FONT = Font(bold=True, color="FFFFFF")
    _HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    _LINK_FONT = Font(color="1155CC", underline="single")

    def __init__(self, path: Path, project_root: Path) -> None:
        self._path = path
        self._project_root = project_root

    def write(
        self,
        applications: Sequence[JobApplication],
        questions: Sequence[UnknownQuestion],
        runs: Sequence[ScanRun] = (),
    ) -> None:
        workbook = Workbook()
        applications_sheet = workbook.active
        applications_sheet.title = "Vagas"
        self._prepare(applications_sheet, self._APPLICATION_COLUMNS)
        for application in applications:
            self._append_application(applications_sheet, application)

        questions_sheet = workbook.create_sheet("Perguntas")
        self._prepare(questions_sheet, self._QUESTION_COLUMNS)
        for question in questions:
            self._append_question(questions_sheet, question)

        runs_sheet = workbook.create_sheet("Execuções")
        self._prepare(runs_sheet, self._RUN_COLUMNS)
        for run in runs:
            self._append_run(runs_sheet, run)

        self._save(workbook)

    def _prepare(self, sheet: Worksheet, columns: Mapping[str, int]) -> None:
        sheet.append(list(columns))
        for cell, width in zip(sheet[1], columns.values(), strict=True):
            cell.font = self._HEADER_FONT
            cell.fill = self._HEADER_FILL
            sheet.column_dimensions[cell.column_letter].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    def _append_application(self, sheet: Worksheet, application: JobApplication) -> None:
        resume_file = self._project_root / application.resume_path if application.resume_path else None
        cells = self._append(
            sheet,
            self._APPLICATION_COLUMNS,
            {
                "Data": application.created_at,
                "Empresa": application.company,
                "Vaga": application.title,
                "Local": application.location,
                "Status": self._STATUS_LABELS[application.status],
                "Aderência": application.fit_percentage / 100,
                "Candidatura": self._APPLY_TYPE_LABELS.get(application.apply_type),
                "Candidatos": application.applicants,
                "Publicada em": application.posted_on,
                "Foco": application.focus,
                "Idioma": application.locale,
                "Vaga de suporte": "Sim" if application.is_support_role else "Não",
                "Tem": ", ".join(application.matched_terms),
                "Falta": ", ".join(application.missing_terms),
                "Observação": application.notes,
                "CV": resume_file.name if resume_file else None,
                "Link da vaga": "Abrir" if application.url else None,
                "Atualizado em": application.updated_at,
            },
        )
        cells["Data"].number_format = cells["Atualizado em"].number_format = self._DATE_TIME_FORMAT
        cells["Publicada em"].number_format = self._DATE_FORMAT
        cells["Aderência"].number_format = "0%"
        if resume_file:
            self._link(cells["CV"], Path(os.path.relpath(resume_file, self._path.parent)).as_posix())
        if application.url:
            self._link(cells["Link da vaga"], application.url)

    def _append_question(self, sheet: Worksheet, question: UnknownQuestion) -> None:
        last_job = f"{question.last_company} - {question.last_title}" if question.last_company else None
        cells = self._append(
            sheet,
            self._QUESTION_COLUMNS,
            {
                "Pergunta": question.question,
                "Ocorrências": question.occurrences,
                "Última vaga": last_job,
                "Primeira vez": question.first_seen_at,
                "Última vez": question.last_seen_at,
            },
        )
        cells["Primeira vez"].number_format = cells["Última vez"].number_format = self._DATE_TIME_FORMAT

    def _append_run(self, sheet: Worksheet, run: ScanRun) -> None:
        cells = self._append(
            sheet,
            self._RUN_COLUMNS,
            {
                "Início": run.started_at,
                "Fim": run.finished_at,
                "Situação": self._RUN_STATUS_LABELS[run.status],
                "Motivo": run.reason,
                "Vagas vistas": run.jobs_seen,
                "Vagas registradas": run.jobs_recorded,
            },
        )
        cells["Início"].number_format = cells["Fim"].number_format = self._DATE_TIME_FORMAT

    @staticmethod
    def _append(sheet: Worksheet, columns: Mapping[str, int], values: Mapping[str, object]) -> dict[str, Cell]:
        sheet.append([values[header] for header in columns])
        sheet.auto_filter.ref = sheet.dimensions
        return dict(zip(columns, sheet[sheet.max_row], strict=True))

    def _link(self, cell: Cell, target: str) -> None:
        cell.hyperlink = target
        cell.font = self._LINK_FONT

    def _save(self, workbook: Workbook) -> None:
        self._path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self._path.with_name(f"~{self._path.stem}.tmp.xlsx")
        workbook.save(temporary)
        try:
            os.replace(temporary, self._path)
        except OSError as error:
            temporary.unlink(missing_ok=True)
            raise SpreadsheetLockedError(
                f"Não foi possível atualizar {self._path} (ela está aberta no Excel?). "
                "Os dados estão salvos no banco e a planilha será atualizada na próxima execução."
            ) from error
